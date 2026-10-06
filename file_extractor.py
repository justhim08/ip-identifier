"""Extract text and indicators from common local document formats."""

import csv
import email
import html
import ipaddress
import json
import re
import zipfile
from email import policy
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit
from xml.etree import ElementTree


SUPPORTED_FORMATS = {
	".csv",
	".docx",
	".eml",
	".html",
	".htm",
	".json",
	".log",
	".md",
	".pdf",
	".rtf",
	".txt",
	".xml",
}
URL_PATTERN = re.compile(r"https?://[^\s<>\"']+", re.IGNORECASE)
EMAIL_PATTERN = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
DOMAIN_PATTERN = re.compile(
	r"(?<![\w.-])(?:[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?\.)+[A-Z]{2,63}(?![\w.-])",
	re.IGNORECASE,
)
IP_CANDIDATE_PATTERN = re.compile(r"(?<![A-Z0-9])(?:[0-9A-F]{0,4}:){2,}[0-9A-F:.]+|(?<![\d.])(?:\d{1,3}\.){3}\d{1,3}", re.IGNORECASE)


class _VisibleTextParser(HTMLParser):
	def __init__(self):
		super().__init__(convert_charrefs=True)
		self.parts = []
		self.hidden_tags = []

	def handle_starttag(self, tag, _attrs):
		tag = tag.lower()
		attributes = dict(_attrs)
		if (
			tag in {"script", "style"}
			or "hidden" in attributes
			or str(attributes.get("aria-hidden", "")).lower() == "true"
		):
			self.hidden_tags.append(tag)

	def handle_endtag(self, tag):
		tag = tag.lower()
		for index in range(len(self.hidden_tags) - 1, -1, -1):
			if self.hidden_tags[index] == tag:
				del self.hidden_tags[index]
				break

	def handle_data(self, data):
		if not self.hidden_tags:
			self.parts.append(data)


def _read_text(path):
	try:
		return path.read_text(encoding="utf-8-sig"), "utf-8"
	except UnicodeDecodeError:
		return path.read_text(encoding="cp1252"), "cp1252"


def _extract_pdf(path):
	try:
		from pypdf import PdfReader
	except ImportError as exc:
		raise RuntimeError("PDF support requires the project dependencies. Install them with: pip install -r requirements.txt") from exc

	from pypdf.errors import PdfReadError

	try:
		reader = PdfReader(str(path))
		text = "\n".join(page.extract_text() or "" for page in reader.pages)
		metadata = {}
		if reader.metadata:
			metadata = {str(key): str(value) for key, value in reader.metadata.items() if value is not None}
	except PdfReadError as exc:
		raise ValueError(f"Could not read PDF file: {exc}") from exc
	return text, metadata


def _extract_docx(path):
	try:
		from docx import Document
		from docx.opc.exceptions import PackageNotFoundError
	except ImportError as exc:
		raise RuntimeError("DOCX support requires the project dependencies. Install them with: pip install -r requirements.txt") from exc

	try:
		document = Document(str(path))
	except (PackageNotFoundError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
		raise ValueError(f"Could not read DOCX file: {exc}") from exc
	parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
	for table in document.tables:
		for row in table.rows:
			parts.append("\t".join(cell.text for cell in row.cells))
	properties = document.core_properties
	metadata = {
		"author": properties.author,
		"title": properties.title,
		"subject": properties.subject,
		"created": properties.created.isoformat() if properties.created else None,
		"modified": properties.modified.isoformat() if properties.modified else None,
	}
	return "\n".join(parts), {key: value for key, value in metadata.items() if value}


def _extract_email(path):
	with path.open("rb") as message_file:
		message = email.message_from_binary_file(message_file, policy=policy.default)
	parts = []
	if message.is_multipart():
		for part in message.walk():
			if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
				content = part.get_content()
				if isinstance(content, str):
					parts.append(content)
	else:
		content = message.get_content()
		if isinstance(content, str):
			parts.append(content)
	metadata = {
		key: str(message[key])
		for key in ("subject", "from", "to", "date")
		if message[key]
	}
	return "\n".join(parts), metadata


def _extract_text(path):
	suffix = path.suffix.lower()
	if suffix == ".pdf":
		return _extract_pdf(path)
	if suffix == ".docx":
		return _extract_docx(path)
	if suffix == ".eml":
		return _extract_email(path)

	text, encoding = _read_text(path)
	metadata = {"encoding": encoding}
	if suffix in {".html", ".htm"}:
		parser = _VisibleTextParser()
		parser.feed(text)
		text = " ".join(parser.parts)
	elif suffix == ".csv":
		rows = csv.reader(text.splitlines())
		text = "\n".join("\t".join(row) for row in rows)
	elif suffix == ".json":
		text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
	elif suffix == ".rtf":
		text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
		text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
		text = text.replace("{", "").replace("}", "")
		text = html.unescape(text)
	return text, metadata


def _extract_indicators(text):
	urls = {match.rstrip(".,;:!?)]}") for match in URL_PATTERN.findall(text)}
	emails = set(EMAIL_PATTERN.findall(text))
	addresses = set()
	for candidate in IP_CANDIDATE_PATTERN.findall(text):
		try:
			addresses.add(str(ipaddress.ip_address(candidate.rstrip("."))))
		except ValueError:
			continue

	domains = set()
	for domain in DOMAIN_PATTERN.findall(text):
		try:
			ipaddress.ip_address(domain)
		except ValueError:
			domains.add(domain.lower())
	domains.update(address.rsplit("@", 1)[1].lower() for address in emails)
	for url in urls:
		try:
			host = urlsplit(url).hostname
		except ValueError:
			host = None
		if host:
			try:
				ipaddress.ip_address(host)
			except ValueError:
				domains.add(host.lower())
	return {
		"urls": sorted(urls),
		"domains": sorted(domains),
		"ip_addresses": sorted(addresses),
		"email_addresses": sorted(emails),
	}


def extract_file(filename):
	"""Extract text, basic metadata, and URL/domain/IP/email indicators."""
	path = Path(filename).expanduser()
	if not path.is_file():
		raise ValueError(f"File not found: {path}")
	if path.suffix.lower() not in SUPPORTED_FORMATS:
		formats = ", ".join(sorted(SUPPORTED_FORMATS))
		raise ValueError(f"Unsupported file type '{path.suffix}'. Supported types: {formats}")

	text, metadata = _extract_text(path)
	return {
		"file": str(path.resolve()),
		"format": path.suffix.lower().lstrip("."),
		"metadata": metadata,
		"text": text,
		"indicators": _extract_indicators(text),
	}
