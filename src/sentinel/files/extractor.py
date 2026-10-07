"""Extract text, metadata, and indicators from supported local document formats."""

import csv
import email
import hashlib
import html
import ipaddress
import json
import mimetypes
import re
import zipfile
from email import policy
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit
from xml.etree import ElementTree


SUPPORTED_FORMATS = {
	".csv",
	".docx",
	".eml",
	".html",
	".htm",
	".json",
	".log",
	".md",
	".pdf",
	".rtf",
	".txt",
	".xml",
}
MAX_EXTRACTION_BYTES = 32 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 4096
MAX_DOCX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_PDF_PAGES = 1000
MAX_EXTRACTED_TEXT_CHARS = 5 * 1024 * 1024
FILE_SIGNATURES = (
	(b"%PDF-", "PDF", "application/pdf"),
	(b"\x89PNG\r\n\x1a\n", "PNG image", "image/png"),
	(b"\xff\xd8\xff", "JPEG image", "image/jpeg"),
	(b"GIF87a", "GIF image", "image/gif"),
	(b"GIF89a", "GIF image", "image/gif"),
	(b"MZ", "Windows executable", "application/vnd.microsoft.portable-executable"),
	(b"\x7fELF", "ELF executable", "application/x-executable"),
	(b"{\\rtf", "RTF document", "application/rtf"),
)
EXTENSION_TYPES = {
	".csv": ("CSV text", "text/csv"),
	".docx": ("DOCX document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
	".eml": ("Email message", "message/rfc822"),
	".htm": ("HTML document", "text/html"),
	".html": ("HTML document", "text/html"),
	".json": ("JSON document", "application/json"),
	".log": ("Plain text", "text/plain"),
	".md": ("Markdown document", "text/markdown"),
	".pdf": ("PDF", "application/pdf"),
	".rtf": ("RTF document", "application/rtf"),
	".txt": ("Plain text", "text/plain"),
	".xml": ("XML document", "application/xml"),
}
URL_PATTERN = re.compile(r"https?://[^\s<>\"']+", re.IGNORECASE)
EMAIL_PATTERN = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
DOMAIN_PATTERN = re.compile(
	r"(?<![\w.-])(?:[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?\.)+[A-Z]{2,63}(?![\w.-])",
	re.IGNORECASE,
)
IP_CANDIDATE_PATTERN = re.compile(r"(?<![A-Z0-9])(?:[0-9A-F]{0,4}:){2,}[0-9A-F:.]+|(?<![\d.])(?:\d{1,3}\.){3}\d{1,3}", re.IGNORECASE)
HASH_PATTERN = re.compile(r"(?<![A-Fa-f0-9])(?:[A-Fa-f0-9]{64}|[A-Fa-f0-9]{40}|[A-Fa-f0-9]{32})(?![A-Fa-f0-9])")
PATH_PATTERN = re.compile(
	r"(?<![\w])(?:[A-Za-z]:\\(?:[^\s<>:\"|?*]+\\)*[^\s<>:\"|?*]+|/(?:[^\s<>:\"|?*]+/)*[^\s<>:\"|?*]+)",
)
FILENAME_PATTERN = re.compile(
	r"(?<![\w.-])[\w.-]+\.(?:exe|dll|sys|bat|cmd|ps1|vbs|js|jar|msi|scr|com|sh|py|docm|xlsm)(?![\w.-])",
	re.IGNORECASE,
)
COMMAND_REFERENCE_PATTERN = re.compile(
	r"\b(?:powershell(?:\.exe)?|cmd\.exe|wscript(?:\.exe)?|cscript(?:\.exe)?|"
	r"rundll32(?:\.exe)?|regsvr32(?:\.exe)?|curl(?:\.exe)?|wget(?:\.exe)?|"
	r"bash|sh|python(?:[23](?:\.exe)?)?)\b",
	re.IGNORECASE,
)


class _VisibleTextParser(HTMLParser):
	def __init__(self):
		super().__init__(convert_charrefs=True)
		self.parts = []
		self.hidden_tags = []
		self.metadata = {}
		self.references = set()
		self.in_title = False

	def handle_starttag(self, tag, _attrs):
		tag = tag.lower()
		attributes = dict(_attrs)
		if tag == "title":
			self.in_title = True
		if tag == "meta":
			name = attributes.get("name") or attributes.get("property")
			content = attributes.get("content")
			if name and content:
				self.metadata[str(name).lower()] = str(content)
		if tag in {"a", "area"} and attributes.get("href"):
			self.references.add(str(attributes["href"]).strip())
		if tag in {"img", "iframe", "script", "link", "source", "video", "audio"}:
			value = attributes.get("src") or attributes.get("href")
			if value:
				self.references.add(str(value).strip())
		if (
			tag in {"script", "style"}
			or "hidden" in attributes
			or str(attributes.get("aria-hidden", "")).lower() == "true"
		):
			self.hidden_tags.append(tag)

	def handle_endtag(self, tag):
		tag = tag.lower()
		if tag == "title":
			self.in_title = False
		for index in range(len(self.hidden_tags) - 1, -1, -1):
			if self.hidden_tags[index] == tag:
				del self.hidden_tags[index]
				break

	def handle_data(self, data):
		if self.in_title:
			title = data.strip()
			if title:
				self.metadata.setdefault("title", title)
		if not self.hidden_tags:
			self.parts.append(data)


def identify_file(filename):
	"""Identify a local file from its signature and report extension separately."""
	path = Path(filename).expanduser()
	if not path.is_file():
		raise ValueError(f"File not found: {path}")
	extension = path.suffix.lower()
	with path.open("rb") as file_handle:
		header = file_handle.read(8192)
	size = path.stat().st_size
	detected_type = None
	detected_mime = None
	confidence = "UNKNOWN"
	for signature, file_type, mime_type in FILE_SIGNATURES:
		if header.startswith(signature):
			detected_type = file_type
			detected_mime = mime_type
			confidence = "HIGH"
			break
	if detected_type is None and size <= MAX_EXTRACTION_BYTES and zipfile.is_zipfile(path):
		with zipfile.ZipFile(path) as archive:
			if "word/document.xml" in archive.namelist():
				detected_type = "DOCX document"
				detected_mime = EXTENSION_TYPES[".docx"][1]
				confidence = "HIGH"
			else:
				detected_type = "ZIP archive"
				detected_mime = "application/zip"
				confidence = "HIGH"
	if detected_type is None:
		sniffed = header.lstrip(b"\xef\xbb\xbf\x00\t\r\n ").lower()
		if sniffed.startswith((b"<!doctype html", b"<html", b"<head", b"<body")):
			detected_type, detected_mime, confidence = "HTML document", "text/html", "MEDIUM"
		elif sniffed.startswith(b"<?xml"):
			detected_type, detected_mime, confidence = "XML document", "application/xml", "MEDIUM"
		elif sniffed.startswith((b"from:", b"to:", b"subject:", b"date:")) and b"\n" in sniffed:
			detected_type, detected_mime, confidence = "Email message", "message/rfc822", "MEDIUM"
	if detected_type is None:
		detected_type, detected_mime = EXTENSION_TYPES.get(extension, ("Unknown", None))
		if detected_type != "Unknown":
			confidence = "LOW"
	if detected_mime is None:
		detected_mime = mimetypes.guess_type(path.name)[0]
	extension_type = EXTENSION_TYPES.get(extension, (None, None))[0]
	return {
		"path": str(path.resolve()),
		"name": path.name,
		"extension": extension,
		"extension_type": extension_type,
		"detected_type": detected_type,
		"mime_type": detected_mime,
		"size": size,
		"readable": True,
		"confidence": confidence,
		"type_mismatch": bool(extension_type and detected_type not in {extension_type, "Unknown"}),
		"evidence": [{
			"source": "file signature" if confidence == "HIGH" else "content sniff" if confidence == "MEDIUM" else "filename extension",
			"value": detected_type,
		}],
	}


def hash_file(filename, algorithms=("sha256", "sha1", "md5"), chunk_size=1024 * 1024):
	"""Calculate local file identification hashes in bounded memory."""
	if chunk_size <= 0:
		raise ValueError("Hash chunk size must be greater than zero.")
	hashers = {}
	for algorithm in algorithms:
		try:
			hashers[algorithm] = hashlib.new(algorithm)
		except ValueError as exc:
			raise ValueError(f"Unsupported hash algorithm: {algorithm}") from exc
	with Path(filename).open("rb") as file_handle:
		for chunk in iter(lambda: file_handle.read(chunk_size), b""):
			for hasher in hashers.values():
				hasher.update(chunk)
	return {name: hasher.hexdigest() for name, hasher in hashers.items()}


def _read_text(path):
	try:
		return path.read_text(encoding="utf-8-sig"), "utf-8"
	except UnicodeDecodeError:
		return path.read_text(encoding="cp1252"), "cp1252"


def _extract_pdf(path):
	try:
		from pypdf import PdfReader
	except ImportError as exc:
		raise RuntimeError("PDF support requires the project dependencies. Install them with: pip install -r requirements.txt") from exc

	from pypdf.errors import PdfReadError

	try:
		reader = PdfReader(str(path))
		parts = []
		characters = 0
		limited = len(reader.pages) > MAX_PDF_PAGES
		for page in reader.pages[:MAX_PDF_PAGES]:
			page_text = page.extract_text() or ""
			remaining = MAX_EXTRACTED_TEXT_CHARS - characters
			if remaining <= 0:
				limited = True
				break
			if len(page_text) > remaining:
				page_text = page_text[:remaining]
				limited = True
			parts.append(page_text)
			characters += len(page_text)
			if limited and characters >= MAX_EXTRACTED_TEXT_CHARS:
				break
		text = "\n".join(parts)
		metadata = {}
		if reader.metadata:
			metadata = {str(key): str(value) for key, value in reader.metadata.items() if value is not None}
		if limited:
			metadata["_extraction_note"] = (
				f"PDF text extraction was limited to {MAX_PDF_PAGES} pages and "
				f"{MAX_EXTRACTED_TEXT_CHARS} characters."
			)
	except PdfReadError as exc:
		raise ValueError(f"Could not read PDF file: {exc}") from exc
	return text, metadata


def _extract_docx(path):
	try:
		from docx import Document
		from docx.opc.exceptions import PackageNotFoundError
	except ImportError as exc:
		raise RuntimeError("DOCX support requires the project dependencies. Install them with: pip install -r requirements.txt") from exc

	try:
		with zipfile.ZipFile(path) as package:
			entries = package.infolist()
			if len(entries) > MAX_ARCHIVE_ENTRIES:
				raise ValueError("DOCX package contains too many archive entries.")
			uncompressed_size = sum(entry.file_size for entry in entries)
			if uncompressed_size > MAX_DOCX_UNCOMPRESSED_BYTES:
				raise ValueError("DOCX package exceeds the safe uncompressed-size limit.")
		document = Document(str(path))
	except (PackageNotFoundError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
		raise ValueError(f"Could not read DOCX file: {exc}") from exc
	parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
	for table in document.tables:
		for row in table.rows:
			parts.append("\t".join(cell.text for cell in row.cells))
	properties = document.core_properties
	hyperlinks = sorted({
		relation.target_ref
		for relation in document.part.rels.values()
		if relation.reltype.endswith("/hyperlink") and relation.target_ref
	})
	metadata = {
		"author": properties.author,
		"title": properties.title,
		"subject": properties.subject,
		"created": properties.created.isoformat() if properties.created else None,
		"modified": properties.modified.isoformat() if properties.modified else None,
		"hyperlinks": hyperlinks,
	}
	return "\n".join(parts), {key: value for key, value in metadata.items() if value}


def _extract_email(path):
	with path.open("rb") as message_file:
		message = email.message_from_binary_file(message_file, policy=policy.default)
	parts = []
	if message.is_multipart():
		for part in message.walk():
			if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
				content = part.get_content()
				if isinstance(content, str):
					parts.append(content)
	else:
		content = message.get_content()
		if isinstance(content, str):
			parts.append(content)
	metadata = {
		key: str(message[key])
		for key in ("subject", "from", "to", "cc", "date", "message-id")
		if message[key]
	}
	attachments = []
	if message.is_multipart():
		for part in message.walk():
			if part.get_content_disposition() != "attachment":
				continue
			payload = part.get_payload(decode=False)
			attachments.append({
				"filename": part.get_filename(),
				"content_type": part.get_content_type(),
				"disposition": part.get_content_disposition(),
				"encoded_size": len(payload) if isinstance(payload, str) else None,
				"payload_opened": False,
			})
	if attachments:
		metadata["attachments"] = attachments
	return "\n".join(parts), metadata


def _extract_text(path):
	suffix = path.suffix.lower()
	if suffix == ".pdf":
		return _extract_pdf(path)
	if suffix == ".docx":
		return _extract_docx(path)
	if suffix == ".eml":
		return _extract_email(path)

	text, encoding = _read_text(path)
	metadata = {"encoding": encoding}
	if suffix in {".html", ".htm"}:
		parser = _VisibleTextParser()
		parser.feed(text)
		text = " ".join(parser.parts)
	elif suffix == ".csv":
		rows = csv.reader(text.splitlines())
		text = "\n".join("\t".join(row) for row in rows)
	elif suffix == ".json":
		text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
	elif suffix == ".rtf":
		text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
		text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
		text = text.replace("{", "").replace("}", "")
		text = html.unescape(text)
	if suffix in {".html", ".htm"}:
		metadata = {
			**metadata,
			"html_metadata": parser.metadata,
			"references": sorted(parser.references),
		}
	return text, metadata


def _extract_indicators(text):
	records = _extract_indicator_records(text)
	values = {
		category: sorted({item["normalized_value"] for item in records.get(category, [])})
		for category in ("url", "domain", "ip_address", "email")
	}
	return {
		"urls": values["url"],
		"domains": values["domain"],
		"ip_addresses": values["ip_address"],
		"email_addresses": values["email"],
	}


def _normalize_url(value):
	try:
		parts = urlsplit(value)
		if not parts.scheme or not parts.hostname:
			return value
		host = parts.hostname.lower()
		if ":" in host and not host.startswith("["):
			host = f"[{host}]"
		port = parts.port
		netloc = host
		if port is not None and not (
			(parts.scheme.lower() == "http" and port == 80)
			or (parts.scheme.lower() == "https" and port == 443)
		):
			netloc = f"{netloc}:{port}"
		return parts._replace(scheme=parts.scheme.lower(), netloc=netloc).geturl()
	except ValueError:
		return value


def _extract_indicator_records(text, metadata=None):
	"""Return normalized, de-duplicated indicators with occurrence evidence."""
	metadata = metadata or {}
	found = {}

	def add(category, original, start, end, source="extracted_text", context_text=None):
		original = original.strip().rstrip(".,;:!?)]}")
		if not original:
			return
		normalized = original
		if category == "url":
			normalized = _normalize_url(original)
		elif category in {"domain", "email"}:
			normalized = original.lower().rstrip(".")
		elif category == "ip_address":
			try:
				normalized = str(ipaddress.ip_address(original.rstrip(".")))
			except ValueError:
				return
		elif category == "hash":
			normalized = original.lower()
		key = (category, normalized)
		evidence_text = context_text if context_text is not None else text
		context_start = max(0, start - 48)
		context_end = min(len(evidence_text), end + 48)
		observation = {
			"source": source,
			"offset": start,
			"context": evidence_text[context_start:context_end].replace("\r", " ").replace("\n", " "),
		}
		item = found.setdefault(key, {
			"value": normalized,
			"normalized_value": normalized,
			"original_values": [],
			"occurrence_count": 0,
			"evidence": [],
		})
		item["occurrence_count"] += 1
		if original not in item["original_values"]:
			item["original_values"].append(original)
		if len(item["evidence"]) < 10 and observation not in item["evidence"]:
			item["evidence"].append(observation)

	for match in URL_PATTERN.finditer(text):
		add("url", match.group(), *match.span())
	for match in EMAIL_PATTERN.finditer(text):
		add("email", match.group(), *match.span())
	for match in IP_CANDIDATE_PATTERN.finditer(text):
		add("ip_address", match.group(), *match.span())
	for match in DOMAIN_PATTERN.finditer(text):
		try:
			ipaddress.ip_address(match.group())
		except ValueError:
			add("domain", match.group(), *match.span())
	for match in HASH_PATTERN.finditer(text):
		add("hash", match.group(), *match.span())
	for match in PATH_PATTERN.finditer(text):
		add("path", match.group(), *match.span())
	for match in FILENAME_PATTERN.finditer(text):
		add("filename", match.group(), *match.span())
	for match in COMMAND_REFERENCE_PATTERN.finditer(text):
		add("command_reference", match.group(), *match.span())

	def scan_metadata(value, source):
		if isinstance(value, dict):
			for key, child in value.items():
				scan_metadata(child, f"{source}.{key}")
		elif isinstance(value, list):
			for index, child in enumerate(value):
				scan_metadata(child, f"{source}[{index}]")
		elif isinstance(value, str):
			patterns = (
				("url", URL_PATTERN),
				("email", EMAIL_PATTERN),
				("ip_address", IP_CANDIDATE_PATTERN),
				("domain", DOMAIN_PATTERN),
				("hash", HASH_PATTERN),
				("path", PATH_PATTERN),
				("filename", FILENAME_PATTERN),
				("command_reference", COMMAND_REFERENCE_PATTERN),
			)
			for category, pattern in patterns:
				for match in pattern.finditer(value):
					if category == "ip_address":
						try:
							ipaddress.ip_address(match.group().rstrip("."))
						except ValueError:
							continue
					elif category == "domain":
						try:
							ipaddress.ip_address(match.group())
						except ValueError:
							pass
						else:
							continue
					add(category, match.group(), *match.span(), source, value)

	scan_metadata(metadata, "metadata")

	by_category = {}
	for key, item in found.items():
		category = key[0]
		by_category.setdefault(category, []).append(item)
	for items in by_category.values():
		items.sort(key=lambda item: item["normalized_value"])
	return by_category


def extract_file(filename):
	"""Extract text, basic metadata, and URL/domain/IP/email indicators."""
	path = Path(filename).expanduser()
	if not path.is_file():
		raise ValueError(f"File not found: {path}")
	if path.suffix.lower() not in SUPPORTED_FORMATS:
		formats = ", ".join(sorted(SUPPORTED_FORMATS))
		raise ValueError(f"Unsupported file type '{path.suffix}'. Supported types: {formats}")

	text, metadata = _extract_text(path)
	return {
		"file": str(path.resolve()),
		"format": path.suffix.lower().lstrip("."),
		"metadata": metadata,
		"text": text,
		"indicators": _extract_indicators(text),
	}
