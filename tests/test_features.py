import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from sentinel.files.extractor import _extract_indicators, extract_file
from sentinel.network.inventory import _parse_neighbors, collect_local_inventory
from sentinel.network.scanner import parse_ports
from sentinel.network.services import identify_service
from sentinel.network.wifi import (
	discover_wifi_access_points,
	parse_windows_networks,
	require_authorization_confirmation,
)
from sentinel.osint.intelligence import lookup_osint, normalize_target
from sentinel.reports.renderer import export_report, render_report, resolve_format
from sentinel.router.audit import ROUTER_TCP_PORTS, audit_router, _validate_router_target


class OsintTests(unittest.TestCase):
	def test_normalize_url_and_domain(self):
		self.assertEqual(normalize_target("https://Example.COM/path?q=1"), ("example.com", "domain"))

	def test_reject_non_public_ip(self):
		with self.assertRaisesRegex(ValueError, "globally routable"):
			normalize_target("127.0.0.1")

	def test_reject_scoped_ipv6(self):
		with self.assertRaisesRegex(ValueError, "Scoped IPv6"):
			normalize_target("https://[2001:4860:4860::8888%25eth0]/")

	def test_domain_lookup_contains_sources(self):
		with patch("sentinel.osint.intelligence._rdap_lookup", return_value={"name": "Example"}), patch(
			"sentinel.osint.intelligence._dns_lookup", return_value={"status": 0, "answers": []}
		), patch("sentinel.osint.intelligence._certificate_names", return_value={"count": 0, "names": [], "truncated": False}):
			report = lookup_osint("example.com")

		self.assertEqual(report["type"], "domain")
		self.assertEqual(set(report["sources"]["dns"]), {"A", "AAAA", "CNAME", "MX", "NS", "TXT", "CAA"})
		self.assertIn("certificate_transparency", report["sources"])

	def test_ip_lookup_includes_passive_ip_intelligence(self):
		with patch("sentinel.osint.intelligence._rdap_lookup", return_value={"name": "Example"}), patch(
			"sentinel.osint.intelligence._dns_lookup", return_value={"status": 0, "answers": []}
		), patch("sentinel.osint.intelligence._ip_intelligence", return_value={"ip": "8.8.8.8", "asn": 15169}):
			report = lookup_osint("8.8.8.8")

		self.assertIn("reverse_dns", report["sources"])
		self.assertEqual(report["sources"]["ip_intelligence"]["asn"], 15169)


class FileExtractionTests(unittest.TestCase):
	def test_extract_indicators_from_text(self):
		indicators = _extract_indicators(
			"Visit https://example.com/path, email analyst@example.org; "
			"IPs 192.0.2.1 and 2001:db8::1."
		)
		self.assertEqual(indicators["urls"], ["https://example.com/path"])
		self.assertIn("example.com", indicators["domains"])
		self.assertIn("example.org", indicators["domains"])
		self.assertEqual(indicators["ip_addresses"], ["192.0.2.1", "2001:db8::1"])
		self.assertEqual(indicators["email_addresses"], ["analyst@example.org"])

	def test_extract_csv_content(self):
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "indicators.csv"
			path.write_text("target,ip\nhttps://example.com,198.51.100.4\n", encoding="utf-8")
			result = extract_file(path)

		self.assertIn("https://example.com", result["text"])
		self.assertIn("198.51.100.4", result["indicators"]["ip_addresses"])

	def test_extract_html_visible_text_only(self):
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "page.html"
			path.write_text(
				"<p>Contact web@example.org</p><script>secret@example.org</script>"
				"<div hidden>hidden text</div>",
				encoding="utf-8",
			)
			result = extract_file(path)

		self.assertIn("web@example.org", result["text"])
		self.assertNotIn("secret@example.org", result["text"])
		self.assertNotIn("hidden text", result["text"])

	def test_extract_docx_text_and_metadata(self):
		from docx import Document

		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "report.docx"
			document = Document()
			document.add_paragraph("Investigate https://example.net")
			document.add_paragraph("Contact ops@example.net")
			document.save(path)
			result = extract_file(path)

		self.assertIn("https://example.net", result["text"])
		self.assertIn("ops@example.net", result["indicators"]["email_addresses"])

	def test_extract_pdf_file(self):
		from pypdf import PdfWriter

		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "report.pdf"
			writer = PdfWriter()
			writer.add_blank_page(width=72, height=72)
			with path.open("wb") as pdf_file:
				writer.write(pdf_file)
			result = extract_file(path)

		self.assertEqual(result["format"], "pdf")
		self.assertEqual(result["text"], "")
		json.dumps(result)

	def test_invalid_pdf_reports_readable_error(self):
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "broken.pdf"
			path.write_text("not a PDF", encoding="utf-8")
			with self.assertRaisesRegex(ValueError, "Could not read PDF"):
				extract_file(path)


	def test_invalid_docx_reports_readable_error(self):
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "broken.docx"
			path.write_text("not a DOCX", encoding="utf-8")
			with self.assertRaisesRegex(ValueError, "Could not read DOCX"):
				extract_file(path)


class RouterAuditTests(unittest.TestCase):
	def test_rejects_public_router_target(self):
		with self.assertRaisesRegex(ValueError, "limited to private"):
			_validate_router_target("8.8.8.8")

	def test_rejects_documentation_ip_as_router_target(self):
		with self.assertRaisesRegex(ValueError, "limited to private"):
			_validate_router_target("192.0.2.10")

	def test_rejects_router_outside_supplied_subnet(self):
		with self.assertRaisesRegex(ValueError, "belong to the supplied subnet"):
			_validate_router_target("192.168.1.1", "192.168.2.0/24")

	def test_audit_scans_only_fixed_tcp_ports_and_collects_advisories(self):
		with patch("sentinel.router.audit._probe_tcp", return_value="closed") as probe, patch(
			"sentinel.router.audit._nvd_advisories", return_value={"status": "candidates", "results": []}
		):
			report = audit_router("192.168.1.1", "Example Router", "1.2.3")

		self.assertEqual(set(report["services"]), {str(port) for port in ROUTER_TCP_PORTS})
		self.assertEqual(probe.call_count, len(ROUTER_TCP_PORTS))
		self.assertEqual(report["udp_services"]["status"], "not_checked")
		self.assertEqual(report["router_identity"]["firmware_supplied_by_user"], "1.2.3")
		self.assertEqual(report["vulnerability_advisories"]["status"], "candidates")

	def test_http_head_only_runs_for_open_management_ports(self):
		def status_for_port(address, port):
			self.assertEqual(str(address), "192.168.1.1")
			return "open" if port == 80 else "closed"

		with patch("sentinel.router.audit._probe_tcp", side_effect=status_for_port), patch(
			"sentinel.router.audit._http_head", return_value={"status_code": 200}
		) as head, patch("sentinel.router.audit._nvd_advisories", return_value={"status": "not_requested"}):
			report = audit_router("192.168.1.1")

		head.assert_called_once()
		self.assertEqual(head.call_args.args[1:], (80, "http"))
		self.assertIn("http_head", report["services"]["80"])
		self.assertNotIn("http_head", report["services"]["443"])


class WifiDiscoveryTests(unittest.TestCase):
	def test_parse_windows_network_list_with_multiple_access_points(self):
		output = (
			"SSID 1 : Home WiFi\n"
			"    Authentication : WPA2-Personal\n"
			"    Encryption : CCMP\n"
			"    BSSID 1 : 00:11:22:33:44:55\n"
			"         Signal : 87%\n"
			"         Radio type : 802.11ax\n"
			"         Channel : 6\n"
			"    BSSID 2 : 00:11:22:33:44:66\n"
			"         Signal : 42%\n"
			"         Channel : 11\n"
			"SSID 2 :\n"
			"    BSSID 1 : aa:bb:cc:dd:ee:ff\n"
			"         Signal : 15%\n"
		)

		access_points = parse_windows_networks(output)

		self.assertEqual(len(access_points), 3)
		self.assertEqual(access_points[0]["ssid"], "Home WiFi")
		self.assertEqual(access_points[0]["authentication"], "WPA2-Personal")
		self.assertEqual(access_points[0]["bssid"], "00:11:22:33:44:55")
		self.assertEqual(access_points[1]["channel"], "11")
		self.assertEqual(access_points[2]["ssid"], "")

	def test_parse_windows_network_without_bssid(self):
		access_points = parse_windows_networks(
			"SSID 1 : Guest\n    Authentication : WPA3-Personal\n    Encryption : GCMP\n"
		)
		self.assertEqual(access_points, [{
			"ssid": "Guest",
			"authentication": "WPA3-Personal",
			"encryption": "GCMP",
		}])

	def test_requires_exact_authorization_confirmation(self):
		with self.assertRaisesRegex(ValueError, "cancelled"):
			require_authorization_confirmation("yes")
		require_authorization_confirmation("I AM AUTHORIZED")

	def test_discovery_runs_only_the_windows_netsh_query(self):
		result = type("CommandResult", (), {"returncode": 0, "stderr": "", "stdout": "SSID 1 : Lab\n    BSSID 1 : 00:11:22:33:44:55\n"})()
		with patch("sentinel.network.wifi.platform.system", return_value="Windows"), patch(
			"sentinel.network.wifi.subprocess.run", return_value=result
		) as run:
			access_points = discover_wifi_access_points()

		self.assertEqual(access_points[0]["ssid"], "Lab")
		self.assertEqual(run.call_args.args[0], ["netsh", "wlan", "show", "networks", "mode=bssid"])

	def test_discovery_rejects_unsupported_platform(self):
		with patch("sentinel.network.wifi.platform.system", return_value="Linux"):
			with self.assertRaisesRegex(RuntimeError, "requires Windows"):
				discover_wifi_access_points()


class NetworkInventoryTests(unittest.TestCase):
	def test_parse_neighbor_cache_for_ipv4_and_ipv6(self):
		neighbors = _parse_neighbors(
			"  192.168.1.1          aa-bb-cc-dd-ee-ff     dynamic\n"
			"fe80::1 dev wlan0 lladdr 00:11:22:33:44:55 REACHABLE\n"
		)

		self.assertEqual(len(neighbors), 2)
		self.assertEqual(neighbors[0]["ip"], "192.168.1.1")
		self.assertEqual(neighbors[0]["mac"], "aa:bb:cc:dd:ee:ff")
		self.assertEqual(neighbors[1]["ip"], "fe80::1")

	def test_collect_local_inventory_uses_read_only_cache_functions(self):
		with patch("sentinel.network.inventory.platform.system", return_value="Windows"), patch(
			"sentinel.network.inventory._windows_inventory",
			return_value=("192.168.1.1", [{"ip": "192.168.1.20", "mac": None, "state": "dynamic"}]),
		), patch(
			"sentinel.network.inventory._get_local_interfaces",
			return_value=[{
				"name": "Ethernet",
				"addresses": [{"address": "192.168.1.20", "ip_version": 4, "network": "192.168.1.0/24"}],
				"mac": "aa:bb:cc:dd:ee:ff",
				"vendor": None,
			}],
		):
			report = collect_local_inventory()

		self.assertEqual(report["default_gateway"], "192.168.1.1")
		self.assertEqual(len(report["neighbors"]), 1)
		self.assertIn("CACHED LOCAL NEIGHBOR SNAPSHOT", report["scope"])
		self.assertEqual(report["interfaces"][0]["name"], "Ethernet")
		self.assertEqual(report["neighbor_snapshot_type"], "CACHED LOCAL NEIGHBOR SNAPSHOT")


class NetworkServiceTests(unittest.TestCase):
	def test_known_service_uses_local_port_mapping(self):
		self.assertEqual(identify_service("192.168.1.1", 22, "closed")["name"], "SSH")

	def test_open_http_service_uses_header_only_identification(self):
		with patch("sentinel.network.services._http_head", return_value={"status_code": 200, "headers": {"server": "test"}}) as head:
			service = identify_service("192.168.1.1", 80, "open")

		self.assertEqual(service["name"], "HTTP")
		self.assertEqual(service["http"]["status_code"], 200)
		head.assert_called_once()


class ScannerTests(unittest.TestCase):
	def test_parse_ports_preserves_valid_port_ranges(self):
		self.assertEqual(parse_ports("22,80,8000-8002"), [22, 80, 8000, 8001, 8002])


class ReportTests(unittest.TestCase):
	def test_renders_json_csv_and_human_readable(self):
		report = {"target": "192.168.1.1", "services": {"80": {"name": "HTTP"}}}

		self.assertIn('"target": "192.168.1.1"', render_report(report, "json"))
		csv_report = render_report(report, "csv")
		self.assertIn("services.80.name,HTTP", csv_report)
		self.assertIn("SENTINEL", render_report(report, "txt"))
		self.assertEqual(resolve_format("report.csv"), "csv")
		self.assertEqual(resolve_format(None, "txt"), "txt")

	def test_exports_to_requested_file(self):
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "report.txt"
			text = export_report({"status": "ok"}, path)

			self.assertEqual(path.read_text(encoding="utf-8"), text)
			self.assertIn("Status", text)


if __name__ == "__main__":
	unittest.main()
